import os


def extract_text_from_docx(file_path: str) -> str:
    """
    Extracts text from a DOCX file using python-docx.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        from docx import Document
        document = Document(file_path)

        paragraphs = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract text from any tables in the resume
        for table in document.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        return "\n".join(paragraphs).strip()

    except ImportError:
        # Fallback reading docx xml if python-docx not installed
        import zipfile
        import xml.etree.ElementTree as ET
        try:
            with zipfile.ZipFile(file_path) as z:
                xml_content = z.read("word/document.xml")
                root = ET.fromstring(xml_content)
                text_elements = root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
                return "".join([node.text for node in text_elements if node.text])
        except Exception as e:
            raise ValueError(f"Failed to read docx document: {str(e)}")
