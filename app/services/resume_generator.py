import os
from pathlib import Path
from app.models.resume import ResumeData
from app.models.analysis import TailoredResumeResult


def generate_docx_resume(
    resume: ResumeData,
    tailored: TailoredResumeResult,
    output_path: str
) -> str:
    """
    Generates an ATS-friendly, single-column Microsoft Word (.docx) resume.
    """
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Set 0.5 inch margins for ATS standards
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        # Candidate Name
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_p.add_run(resume.full_name)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(16)

        # Contact info
        contacts = []
        if resume.email:
            contacts.append(resume.email)
        if resume.phone:
            contacts.append(resume.phone)
        if resume.location:
            contacts.append(resume.location)
        if resume.linkedin:
            contacts.append(resume.linkedin)
        if resume.github:
            contacts.append(resume.github)

        if contacts:
            contact_p = doc.add_paragraph()
            contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c_run = contact_p.add_run("  •  ".join(contacts))
            c_run.font.name = "Arial"
            c_run.font.size = Pt(10)
            c_run.font.color.rgb = RGBColor(90, 90, 90)

        # Professional Summary
        doc.add_heading("PROFESSIONAL SUMMARY", level=1)
        sum_p = doc.add_paragraph(tailored.tailored_summary or resume.summary)
        sum_p.paragraph_format.space_after = Pt(8)

        # Technical Skills
        doc.add_heading("TECHNICAL & PROFESSIONAL SKILLS", level=1)
        for cat_name, skill_list in tailored.prioritized_skills.items():
            sp = doc.add_paragraph()
            r_cat = sp.add_run(f"{cat_name}: ")
            r_cat.bold = True
            sp.add_run(", ".join(skill_list))
            sp.paragraph_format.space_after = Pt(3)

        # Professional Experience
        doc.add_heading("WORK EXPERIENCE", level=1)
        for exp in tailored.tailored_experience:
            role = exp.get("role", "")
            company = exp.get("company", "")
            dates = exp.get("start_date", "")
            if exp.get("end_date"):
                dates += f" - {exp.get('end_date')}"

            exp_p = doc.add_paragraph()
            r_role = exp_p.add_run(role)
            r_role.bold = True
            r_comp = exp_p.add_run(f" | {company}")
            r_comp.italic = True
            if dates:
                r_dates = exp_p.add_run(f" ({dates})")
                r_dates.font.color.rgb = RGBColor(100, 100, 100)

            exp_p.paragraph_format.space_after = Pt(2)

            for b in exp.get("bullets", []):
                bp = doc.add_paragraph(b, style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)

        # Key Projects
        if resume.projects:
            doc.add_heading("KEY PROJECTS", level=1)
            for proj in resume.projects:
                pp = doc.add_paragraph()
                r_proj = pp.add_run(proj.title)
                r_proj.bold = True
                if proj.technologies:
                    r_tech = pp.add_run(f" ({', '.join(proj.technologies)})")
                    r_tech.italic = True
                pp.paragraph_format.space_after = Pt(2)

                for b in proj.bullets:
                    bp = doc.add_paragraph(b, style="List Bullet")
                    bp.paragraph_format.space_after = Pt(2)

        # Education
        if resume.education:
            doc.add_heading("EDUCATION", level=1)
            for edu in resume.education:
                ed_p = doc.add_paragraph()
                degree_text = f"{edu.degree} - {edu.institution}"
                if edu.graduation_year:
                    degree_text += f" ({edu.graduation_year})"
                ed_p.add_run(degree_text)
                ed_p.paragraph_format.space_after = Pt(3)

        # Certifications
        if resume.certifications:
            doc.add_heading("CERTIFICATIONS", level=1)
            cert_p = doc.add_paragraph(", ".join(resume.certifications))
            cert_p.paragraph_format.space_after = Pt(4)

        doc.save(output_path)
        return output_path

    except ImportError:
        # Fallback to plain text document if python-docx not installed
        with open(output_path.replace(".docx", ".txt"), "w", encoding="utf-8") as f:
            f.write(f"{resume.full_name}\n")
            f.write(f"SUMMARY:\n{tailored.tailored_summary}\n\n")
            f.write(f"EXPERIENCE:\n")
            for e in tailored.tailored_experience:
                f.write(f"- {e.get('role')} at {e.get('company')}\n")
                for b in e.get("bullets", []):
                    f.write(f"  * {b}\n")
        return output_path.replace(".docx", ".txt")


def generate_pdf_resume(
    resume: ResumeData,
    tailored: TailoredResumeResult,
    output_path: str
) -> str:
    """
    Generates a PDF resume using ReportLab.
    """
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=16,
            leading=20,
            alignment=1,
            textColor=colors.HexColor('#111827')
        )
        h2_style = ParagraphStyle(
            'H2Style',
            parent=styles['Heading2'],
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#1e40af'),
            spaceBefore=8,
            spaceAfter=4
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#374151')
        )

        story = []
        story.append(Paragraph(resume.full_name, title_style))
        story.append(Spacer(1, 4))

        # Contacts
        contacts = [c for c in [resume.email, resume.phone, resume.location, resume.linkedin] if c]
        if contacts:
            story.append(Paragraph(" • ".join(contacts), ParagraphStyle('Contact', parent=body_style, alignment=1)))
            story.append(Spacer(1, 8))

        # Summary
        story.append(Paragraph("PROFESSIONAL SUMMARY", h2_style))
        story.append(Paragraph(tailored.tailored_summary or resume.summary, body_style))
        story.append(Spacer(1, 6))

        # Skills
        story.append(Paragraph("SKILLS", h2_style))
        for cat, skills in tailored.prioritized_skills.items():
            story.append(Paragraph(f"<b>{cat}:</b> {', '.join(skills)}", body_style))
        story.append(Spacer(1, 6))

        # Experience
        story.append(Paragraph("EXPERIENCE", h2_style))
        for exp in tailored.tailored_experience:
            story.append(Paragraph(f"<b>{exp.get('role')}</b> | <i>{exp.get('company')}</i>", body_style))
            for b in exp.get("bullets", []):
                story.append(Paragraph(f"• {b}", ParagraphStyle('Bullet', parent=body_style, leftIndent=12)))
            story.append(Spacer(1, 4))

        doc.build(story)
        return output_path

    except ImportError:
        # If reportlab is missing, fall back to docx
        docx_path = output_path.replace(".pdf", ".docx")
        return generate_docx_resume(resume, tailored, docx_path)
